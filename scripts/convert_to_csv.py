import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
import csv
import os

# === 1. Convert iPhone Transactions ===
print("Converting iPhone Transactions...")
wb = openpyxl.load_workbook('Iphone Sales Data (1).xlsx', read_only=True)
ws = wb['iPhone Transactions']

with open('iphone_transactions.csv', 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    row_count = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = list(row)
        if i == 0:
            # Header row
            writer.writerow(vals)
        else:
            # Fix Amount column (index 8) — it's a formula like =D2*H2
            # Recalculate as Quantity * Unit_Price
            quantity = vals[3]  # Quantity
            unit_price = vals[7]  # Unit_Price
            if quantity is not None and unit_price is not None:
                vals[8] = round(float(quantity) * float(unit_price), 2)
            else:
                vals[8] = None

            # Format dates as YYYY-MM-DD strings for BigQuery
            for date_idx in [0, 9]:  # Trandate, LastInvDate
                if vals[date_idx] is not None:
                    vals[date_idx] = vals[date_idx].strftime('%Y-%m-%d')

            writer.writerow(vals)
        row_count += 1

wb.close()
csv_size = os.path.getsize('iphone_transactions.csv')
print(f"iphone_transactions.csv: {row_count} rows (incl header), {csv_size:,} bytes")

# === 2. Convert Product Master ===
print("\nConverting Product Master...")
wb2 = openpyxl.load_workbook('Product Master (1).xlsx', read_only=True)
ws2 = wb2['Product Master']

with open('product_master.csv', 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    row_count2 = 0
    for row in ws2.iter_rows(values_only=True):
        writer.writerow(list(row))
        row_count2 += 1

wb2.close()
csv_size2 = os.path.getsize('product_master.csv')
print(f"product_master.csv: {row_count2} rows (incl header), {csv_size2:,} bytes")

# === Quick validation ===
print("\n--- Validation: iphone_transactions.csv ---")
with open('iphone_transactions.csv', 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    header = next(reader)
    print(f"Columns ({len(header)}): {header}")

    total_amount = 0
    total_qty = 0
    data_rows = 0
    for row in reader:
        data_rows += 1
        try:
            total_qty += int(row[3])
            total_amount += float(row[8])
        except (ValueError, IndexError):
            pass

    print(f"Data rows: {data_rows}")
    print(f"Total Quantity: {total_qty:,}")
    print(f"Total Amount (calculated): ${total_amount:,.2f}")

print("\n--- Validation: product_master.csv ---")
with open('product_master.csv', 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    header = next(reader)
    print(f"Columns ({len(header)}): {header}")
    products = 0
    for row in reader:
        products += 1
    print(f"Products: {products}")

print("\n✅ CSV files ready for BigQuery upload!")
